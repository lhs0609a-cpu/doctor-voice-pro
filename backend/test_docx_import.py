"""워드 읽기 검사 — 실제 .docx 를 만들어서 넣고, 서식이 살아남는지 본다.

여기서 확인하는 것: 굵게·기울임·밑줄·글자색·글자크기, 소제목, 인용구, 목록, 표, 문서 안 사진,
그리고 무엇보다 **문서에 놓인 순서**. 표가 조용히 사라지던 것이 이 순회 방식의 문제였다.
"""
import io
import os
import unittest
import zlib

os.environ.setdefault('DATABASE_URL', 'sqlite+aiosqlite:///./test-unused.db')
os.environ.setdefault('DATABASE_URL_SYNC', 'sqlite:///./test-unused.db')

from docx import Document
from docx.shared import Pt, RGBColor

from app.services.docx_import import block_text, flatten_blocks, parse_docx, spans_text


def _png(width: int = 2, height: int = 2) -> bytes:
    """의존성 없이 만드는 최소 PNG."""
    def chunk(tag: bytes, payload: bytes) -> bytes:
        return (len(payload).to_bytes(4, 'big') + tag + payload
                + zlib.crc32(tag + payload).to_bytes(4, 'big'))
    header = width.to_bytes(4, 'big') + height.to_bytes(4, 'big') + bytes([8, 2, 0, 0, 0])
    raw = b''.join(b'\x00' + b'\xff\x00\x00' * width for _ in range(height))
    return (b'\x89PNG\r\n\x1a\n' + chunk(b'IHDR', header)
            + chunk(b'IDAT', zlib.compress(raw)) + chunk(b'IEND', b''))


def _save(document) -> bytes:
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _doc_with_everything() -> bytes:
    document = Document()
    document.add_paragraph('아토피 초기 증상 확인법', style='Heading 1')
    intro = document.add_paragraph()
    intro.add_run('가려움이 ')
    run = intro.add_run('2주 이상')
    run.bold = True
    intro.add_run(' 이어지면 ')
    tail = intro.add_run('진료를 권합니다')
    tail.italic = True
    tail.underline = True
    document.add_paragraph('진료 전 확인할 것', style='Heading 2')
    marked = document.add_paragraph()
    colored = marked.add_run('보습이 가장 중요합니다')
    colored.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
    colored.font.size = Pt(15)
    document.add_paragraph('환자분들이 가장 많이 묻는 질문입니다.', style='Quote')
    document.add_paragraph('아침 세안 후 보습', style='List Bullet')
    document.add_paragraph('저녁 목욕 후 보습', style='List Bullet')
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).paragraphs[0].add_run('단계').bold = True
    table.cell(0, 1).paragraphs[0].add_run('할 일').bold = True
    table.cell(1, 0).text = '1일차'
    table.cell(1, 1).text = '보습제 바르기'
    document.add_paragraph().add_run().add_picture(io.BytesIO(_png()))
    document.add_paragraph('자세한 내용은 진료 때 안내드립니다.')
    return _save(document)


class ParseTests(unittest.TestCase):
    def setUp(self):
        self.parsed = parse_docx(_doc_with_everything(), name='아토피.docx')
        self.kinds = [b['type'] for b in self.parsed.blocks]

    def test_first_heading_becomes_the_title(self):
        self.assertEqual(self.parsed.title, '아토피 초기 증상 확인법')
        self.assertNotIn('아토피 초기 증상 확인법', [block_text(b) for b in self.parsed.blocks])

    def test_document_order_is_kept(self):
        self.assertEqual(self.kinds, ['text', 'heading', 'text', 'quote', 'list', 'table', 'image', 'text'])

    def test_bold_italic_underline_survive(self):
        spans = self.parsed.blocks[0]['spans']
        self.assertEqual(spans_text(spans), '가려움이 2주 이상 이어지면 진료를 권합니다')
        self.assertEqual([s.get('b') for s in spans], [None, True, None, None])
        self.assertTrue(spans[3].get('i') and spans[3].get('u'))

    def test_runs_with_the_same_look_are_joined(self):
        # 워드는 한 문장을 run 여러 개로 쪼개 둔다. 서식이 같으면 하나로 붙어야 한다.
        self.assertEqual(len(self.parsed.blocks[0]['spans']), 4)

    def test_color_and_size_survive(self):
        span = self.parsed.blocks[2]['spans'][0]
        self.assertEqual(span['color'], '#c00000')
        self.assertEqual(span['size'], 15.0)

    def test_heading_level_is_read(self):
        heading = self.parsed.blocks[1]
        self.assertEqual((heading['level'], block_text(heading)), (2, '진료 전 확인할 것'))

    def test_quote_is_its_own_block(self):
        self.assertEqual(block_text(self.parsed.blocks[3]), '환자분들이 가장 많이 묻는 질문입니다.')

    def test_list_items_are_grouped(self):
        listing = self.parsed.blocks[4]
        self.assertEqual([spans_text(i) for i in listing['items']], ['아침 세안 후 보습', '저녁 목욕 후 보습'])
        self.assertFalse(listing['ordered'])

    def test_table_keeps_cells_and_marks_header(self):
        table = self.parsed.blocks[5]
        self.assertTrue(table['header'])
        self.assertEqual([[spans_text(c) for c in row] for row in table['rows']],
                         [['단계', '할 일'], ['1일차', '보습제 바르기']])

    def test_image_becomes_a_data_url(self):
        self.assertTrue(self.parsed.blocks[6]['image'].startswith('data:image/png;base64,'))

    def test_plain_text_includes_table_and_list(self):
        text = self.parsed.text
        self.assertIn('1일차 | 보습제 바르기', text)
        self.assertIn('저녁 목욕 후 보습', text)
        self.assertNotIn('data:image', text)

    def test_flatten_keeps_images_and_merges_text(self):
        flat = flatten_blocks(self.parsed.blocks)
        self.assertEqual([b['type'] for b in flat], ['text', 'image', 'text'])
        self.assertIn('단계 | 할 일', flat[0]['content'])


class TitleTests(unittest.TestCase):
    def test_short_first_line_becomes_the_title(self):
        document = Document()
        document.add_paragraph('무릎 통증 원인 정리')
        document.add_paragraph('본문입니다.')
        parsed = parse_docx(_save(document), name='무릎.docx')
        self.assertEqual(parsed.title, '무릎 통증 원인 정리')
        self.assertEqual(len(parsed.blocks), 1)

    def test_long_first_line_keeps_the_body_and_uses_the_file_name(self):
        document = Document()
        document.add_paragraph('가' * 80)
        parsed = parse_docx(_save(document), name='무릎 통증.docx')
        self.assertEqual(parsed.title, '무릎 통증')
        self.assertEqual(len(parsed.blocks), 1)

    def test_korean_heading_style_name_is_recognized(self):
        document = Document()
        heading = document.add_paragraph('허리 디스크 안내')
        heading.style = document.styles['Heading 1']
        heading.style.name = '제목 1'          # 한국어 워드가 쓰는 이름
        document.add_paragraph('본문입니다.')
        parsed = parse_docx(_save(document), name='x.docx')
        self.assertEqual(parsed.title, '허리 디스크 안내')


class RejectTests(unittest.TestCase):
    def test_broken_file_is_reported(self):
        with self.assertRaises(ValueError) as caught:
            parse_docx(b'not a docx at all', name='깨진.docx')
        self.assertIn('워드 파일을 읽지 못했습니다', str(caught.exception))

    def test_empty_document_is_reported(self):
        with self.assertRaises(ValueError):
            parse_docx(_save(Document()), name='빈.docx')

    def test_pictures_without_text_are_reported(self):
        document = Document()
        document.add_paragraph().add_run().add_picture(io.BytesIO(_png()))
        with self.assertRaises(ValueError) as caught:
            parse_docx(_save(document), name='사진만.docx')
        self.assertIn('글이 없습니다', str(caught.exception))


if __name__ == '__main__':
    unittest.main()
