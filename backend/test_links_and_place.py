"""원고 속 링크와, 아이디마다 저장해 둔 링크·플레이스.

두 가지를 잡아 둔다.
1) 워드에 걸어 둔 하이퍼링크는 글자도 주소도 살아남는다. python-docx 의 paragraph.runs 는
   w:hyperlink 안의 글자를 아예 돌려주지 않아서, 예전에는 '여기서 예약하세요'의 '여기서'가
   통째로 사라졌다(2026-09-22 실측).
2) 블로그(아이디)에 저장해 둔 링크·플레이스는 글 끝에 한 줄씩 붙는다. 네이버는 한 줄짜리
   주소를 링크 카드·지도 카드로 바꿔 주므로, 주소만 제자리에 있으면 된다.
"""
import io
import unittest

import docx
from docx import Document
from docx.oxml.shared import OxmlElement, qn

from app.api.campaign import JobBlock, _with_footer
from app.models.campaign import Blog
from app.services import docx_import


def _doc_with_link(text: str = "여기서", url: str = "https://example.com/reserve") -> bytes:
    document = Document()
    document.add_paragraph("아토피 예약 안내", style="Heading 1")
    paragraph = document.add_paragraph("예약은 ")
    rid = document.part.relate_to(url, docx.opc.constants.RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    node = OxmlElement("w:t")
    node.text = text
    run.append(node)
    link.append(run)
    paragraph._p.append(link)
    paragraph.add_run(" 하시면 됩니다.")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


class ManuscriptLinkTest(unittest.TestCase):
    def test_hyperlink_text_is_not_swallowed(self):
        parsed = docx_import.parse_docx(_doc_with_link(), name="링크.docx")
        self.assertIn("예약은 여기서 하시면 됩니다.", parsed.text)

    def test_the_address_comes_along_as_its_own_line(self):
        """네이버는 한 줄짜리 주소를 링크 카드로 만든다 — 그래야 눌린다."""
        parsed = docx_import.parse_docx(_doc_with_link(), name="링크.docx")
        lines = [docx_import.block_text(b) for b in parsed.blocks]
        self.assertIn("https://example.com/reserve", lines)
        self.assertEqual(parsed.links, [{"text": "여기서", "url": "https://example.com/reserve"}])

    def test_a_visible_address_is_not_repeated(self):
        parsed = docx_import.parse_docx(_doc_with_link(text="https://example.com/reserve"), name="링크.docx")
        lines = [docx_import.block_text(b) for b in parsed.blocks]
        self.assertEqual(sum(l.count("https://example.com/reserve") for l in lines), 1)

    def test_the_link_style_survives_on_the_words(self):
        parsed = docx_import.parse_docx(_doc_with_link(), name="링크.docx")
        spans = [s for b in parsed.blocks for s in b.get("spans", [])]
        linked = [s for s in spans if s.get("href")]
        self.assertEqual([s["t"] for s in linked], ["여기서"])


def _blog(**kwargs) -> Blog:
    return Blog(id="b", user_id="u", client_id="c", blog_id="testblog", **kwargs)


class SavedFooterTest(unittest.TestCase):
    def body(self, blocks):
        return [b.content for b in blocks]

    def test_saved_link_and_place_are_appended_once(self):
        blocks = [JobBlock(type="text", content="본문입니다.")]
        blog = _blog(footer_link_url="https://clinic.example.com/reserve", footer_link_label="예약은 여기서",
                     place_url="https://naver.me/abcd1234", place_label="오시는 길")
        self.assertEqual(self.body(_with_footer(blocks, blog)), [
            "본문입니다.", "예약은 여기서", "https://clinic.example.com/reserve",
            "오시는 길", "https://naver.me/abcd1234",
        ])

    def test_nothing_is_added_when_nothing_is_saved(self):
        blocks = [JobBlock(type="text", content="본문입니다.")]
        self.assertEqual(self.body(_with_footer(blocks, _blog())), ["본문입니다."])

    def test_an_address_already_in_the_manuscript_is_not_added_twice(self):
        """원고가 이미 그 주소를 들고 있으면 글 끝에 또 붙이지 않는다."""
        blocks = [JobBlock(type="text", content="예약: https://clinic.example.com/reserve")]
        blog = _blog(footer_link_url="https://clinic.example.com/reserve", footer_link_label="예약은 여기서")
        self.assertEqual(self.body(_with_footer(blocks, blog)),
                         ["예약: https://clinic.example.com/reserve"])

    def test_the_label_is_optional(self):
        blocks = [JobBlock(type="text", content="본문입니다.")]
        blog = _blog(place_url="https://naver.me/abcd1234")
        self.assertEqual(self.body(_with_footer(blocks, blog)), ["본문입니다.", "https://naver.me/abcd1234"])


if __name__ == "__main__":
    unittest.main()
