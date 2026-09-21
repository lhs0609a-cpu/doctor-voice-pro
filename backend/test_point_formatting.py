import io
import unittest
from copy import deepcopy
from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pydantic import ValidationError
from app.services.point_formatting import PointFormatting, apply_points
from app.services.docx_import import parse_docx, block_text


class PointFormattingTests(unittest.TestCase):
    def test_distinct_treatments_preserve_all_text_and_order(self):
        blocks = [{'type':'text','content':f'중요한 기준 {n}을 먼저 확인해야 합니다.'} for n in range(12)]
        before = deepcopy(blocks)
        output = apply_points(blocks, {'enabled':True, 'phrases':['중요한 기준']})
        self.assertEqual(blocks, before)
        self.assertEqual([block_text(b) or b.get('content') for b in output], [b['content'] for b in blocks])
        self.assertLessEqual(sum(bool(b.get('spans')) for b in output), 6)
        self.assertTrue(any(s.get('b') for b in output for s in b.get('spans', [])))
        self.assertTrue(any(s.get('color') for b in output for s in b.get('spans', [])))
        self.assertTrue(any(s.get('background') for b in output for s in b.get('spans', [])))

    def test_existing_word_styles_images_tables_links_untouched(self):
        blocks = [{'type':'text','spans':[{'t':'중요한 기준','b':True}]},
                  {'type':'image','image':'data:image/png;base64,test'},
                  {'type':'table','rows':[[[{'t':'중요한 기준'}]]]},
                  {'type':'text','content':'중요한 기준 https://example.com'}]
        self.assertEqual(apply_points(blocks, {'enabled':True,'phrases':['중요한 기준']}), blocks)

    def test_disabling_features_does_not_apply_them(self):
        blocks = [{'type':'text','content':'개인별 상태 확인이 중요합니다.'} for _ in range(20)]
        cfg = {'enabled':True,'bold':False,'color':False,'background':False,'quote':True}
        output = apply_points(blocks,cfg)
        self.assertTrue(any(b['type']=='quote' for b in output))
        self.assertFalse(any(s.get('b') or s.get('color') or s.get('background') for b in output for s in b.get('spans',[])))
        self.assertEqual(apply_points(blocks, {'enabled':False}), blocks)

    def test_color_payload_cannot_inject_markup_or_css(self):
        for value in ('red','url(javascript:alert(1))','#fff;display:none'):
            with self.assertRaises(ValidationError):
                PointFormatting(text_color=value)

    def test_word_highlight_and_shading_survive_import(self):
        doc = Document()
        doc.add_heading('서식 검증 원고', level=1)
        p = doc.add_paragraph()
        p.add_run('노란 강조').font.highlight_color = WD_COLOR_INDEX.YELLOW
        run = p.add_run('사용자 배경')
        shade = OxmlElement('w:shd')
        shade.set(qn('w:fill'),'FFF8B2')
        run._r.get_or_add_rPr().append(shade)
        buffer = io.BytesIO(); doc.save(buffer)
        parsed = parse_docx(buffer.getvalue())
        spans = [s for b in parsed.blocks for s in b.get('spans',[])]
        self.assertIn('#ffff00',[s.get('background') for s in spans])
        self.assertIn('#fff8b2',[s.get('background') for s in spans])


if __name__ == '__main__':
    unittest.main()
