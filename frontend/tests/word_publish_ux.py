"""Browser check of Word upload (picker and drag-and-drop), formatting, and interval scheduling.

API responses are isolated fixtures. Real Naver editing is verified separately.
"""
import asyncio
import base64
import io
import os
from pathlib import Path
from urllib.parse import urlparse
from docx import Document
from playwright.async_api import async_playwright, expect

BASE=os.environ.get('UX_BASE_URL', 'http://127.0.0.1:8317').rstrip('/')
OUTPUT=Path(__file__).resolve().parents[2]/'output/naver-formatting-verification'

async def main():
    config=dict(enabled=False,bold=True,quote=True,color=True,background=True,text_color='#0078cb',background_color='#fff8b2',phrases=[])
    blog=dict(id='b',blog_id='fixture_blog',label='테스트 블로그',status='active',daily_limit=2)
    client=dict(id='h',name='테스트 병원',blogs=[blog],briefs=[],diseases=[],treatments=[],regions=[])
    campaign=dict(id='c',name='Word 예약',client_id='h',client_name='테스트 병원',blog_ids=['b'],settings={},stats={},step=3,status='draft')
    draft=dict(id='d',title='중요 포인트 테스트',source='upload',status='ready',char_count=200,checks={},tags=[],image_count_target=0,image_plan=[])
    # 이 블로그에는 이미 예약 한 건이 걸려 있다 — 화면은 그 다음부터를 기본으로 골라야 한다.
    booked=dict(blog_ref_id='b',label='테스트 블로그',count=1,last_at='2026-09-23T10:00',
                scanned_at='2026-09-22T12:00',stale=False,note=None)
    state={'drafts':[],'saves':0,'scheduled':False,'rescan':0,'modes':[],'deleted':[]}
    async def route(r):
        req=r.request; path=urlparse(req.url).path
        if '/api/' not in path:
            if req.url.startswith(BASE):await r.continue_()
            else:await r.abort()
            return
        response={}
        if req.method=='DELETE' and '/drafts/' in path:
            did=path.rsplit('/',1)[-1]
            state['deleted'].append(did)
            state['drafts']=[d for d in state['drafts'] if d['id']!=did]
            response={'success':True}
        elif path.endswith('/campaigns'):response=[campaign]
        elif path.endswith('/clients'):response=[client]
        elif path.endswith('/clients/h'):response=client
        elif path.endswith('/campaigns/c'):response=campaign
        elif path.endswith('/agent/status'):response={'online':False,'running':False,'devices':[]}
        elif path.endswith('/formatting'):
            if req.method=='PUT':config.update(req.post_data_json);state['saves']+=1
            response=config
        elif path.endswith('/drafts/upload'):
            # 이제 한 건씩 따로 올라온다 — 파일 하나가 깨져도 나머지는 살아야 하기 때문이다.
            assert b'word/document.xml' in req.post_data_buffer or b'PK' in req.post_data_buffer
            n=len(state['drafts'])+1
            row=dict(draft,id=f'd{n}',title=f'{draft["title"]} {n}')
            if n==2:   # 병원이 등록한 금칙어에 걸린 원고 한 건
                row=dict(row,status='needs_review',
                         checks={'forbidden':['최고'],'medical_law':[],'ok':False})
            if n==3:   # 올릴 때 서버가 알아서 고친 원고
                row=dict(row,checks={'forbidden':[],'medical_law':[],'ok':True,
                                     'auto_fixed':[{'from':'완치','to':'증상 개선','category':'치료효과_보장'},
                                                   {'from':'검사비는 30,000원입니다.','to':'','category':'가격_할인'}]})
            state['drafts'].append(row);response=[row]
        elif path.endswith('/drafts'):response=state['drafts']
        elif path.endswith('/formatting-preview'):
            assert req.post_data_json['enabled']
            response={'title':draft['title'],'blocks':[{'type':'text','spans':[{'t':'핵심 기준','b':True},{'t':'과 일반 본문입니다.'}]},{'type':'quote','spans':[{'t':'핵심은 개인별 상태에 맞는 선택입니다.'}]}]}
        elif path.endswith('/reservations/rescan'):
            state['rescan']+=1;response={'requested':1}
        elif path.endswith('/reservations'):
            response=[booked]
        elif path.endswith('/schedule/preview') or path.endswith('/schedule/commit'):
            body=req.post_data_json
            assert 'd2' not in body['draft_ids'],body['draft_ids']   # 뺀 원고는 예약에 안 간다
            assert body['mode']=='interval' and body['every_minutes']==120
            assert body['start_at'][:10]==body['start_date']
            state['modes'].append(body['start_mode'])
            response={'total':1,'assigned':[{'draft_id':'d1','title':draft['title'],'blog_ref_id':'b','blog_label':'테스트 블로그','scheduled_at':'2026-09-23T12:00:00'}],
                      'unassigned':0,'calendar':[],'warnings':[],'starts_after':booked['last_at'],'reservations':[booked]}
            if path.endswith('/commit'):state['scheduled']=True
        elif path.endswith('/autopilot'):
            response={'enabled':False,'config':dict(image_count=0,landing_url='',landing_purpose='',daily_posts=2)}
        elif any(path.endswith(suffix) for suffix in ('/keywords','/jobs','/tasks','/devices','/briefs','/collections')):response=[]
        await r.fulfill(json=response)
    async with async_playwright() as pw:
        browser=await pw.chromium.launch(headless=True)
        context=await browser.new_context(viewport={'width':1440,'height':1100})
        await context.add_init_script("localStorage.setItem('access_token','fixture');localStorage.setItem('user',JSON.stringify({id:'u',name:'테스트',email:'test@example.invalid'}));")
        await context.route('**/*',route)
        page=await context.new_page();errors=[]
        page.set_default_timeout(120000)
        page.on('pageerror',lambda e:errors.append(e.stack or str(e)))
        await page.goto(BASE+'/dashboard/one-stop',timeout=120000)
        await page.get_by_role('button',name='Word 원고 올리기').click()
        doc=Document();doc.add_heading(draft['title'],level=1);doc.add_paragraph('핵심 기준과 일반 본문입니다.')
        data=io.BytesIO();doc.save(data)
        await page.get_by_label('예약할 Word 파일').set_input_files({'name':'서식.docx','mimeType':'application/vnd.openxmlformats-officedocument.wordprocessingml.document','buffer':data.getvalue()})
        await expect(page.get_by_role('button',name='선택한 1개 원고 예약하기')).to_be_enabled()
        await expect(page.get_by_text('올린 파일 1개',exact=False)).to_be_visible()

        # 끌어다 놓기 — 두 개를 한 번에. 브라우저가 파일을 열어 버리지 않고 앱이 받아야 한다.
        payload=base64.b64encode(data.getvalue()).decode()
        drop=page.get_by_label('예약할 Word 파일').locator('xpath=..')
        transfer=await page.evaluate_handle('''([b64,names])=>{
            const bin=atob(b64), bytes=new Uint8Array(bin.length);
            for(let i=0;i<bin.length;i++) bytes[i]=bin.charCodeAt(i);
            const dt=new DataTransfer();
            for(const name of names) dt.items.add(new File([bytes],name,
                {type:'application/vnd.openxmlformats-officedocument.wordprocessingml.document'}));
            return dt;
        }''',[payload,['밀양성장클리닉 - 스트레스.docx','거제성장클리닉 - 스트레스.docx']])
        await drop.dispatch_event('drop',{'dataTransfer':transfer})
        await expect(page.get_by_text('올린 파일 3개',exact=False)).to_be_visible()
        await expect(page.get_by_text('밀양성장클리닉 - 스트레스.docx',exact=False)).to_be_visible()
        # 검토 필요 한 건은 고를 수 없고, 이유를 펴면 무엇이 걸렸는지 적혀 있다.
        await expect(page.get_by_role('button',name='선택한 2개 원고 예약하기')).to_be_enabled()
        await page.get_by_role('button',name='검토 필요',exact=False).click()
        await expect(page.get_by_text('병원 금칙어',exact=False)).to_be_visible()
        await expect(page.get_by_text('최고',exact=False).first).to_be_visible()
        # 그대로 빼 버린다 — 목록에서도 사라지고 예약에도 안 간다.
        await page.get_by_role('button',name='이 원고 빼기',exact=True).click()
        await expect(page.get_by_role('button',name='선택한 2개 원고 예약하기')).to_be_enabled()

        # 자동으로 고친 원고는 무엇을 고쳤는지 알려 준다(막지 않는다).
        await page.get_by_role('button',name='표현 2곳 고침',exact=True).click()
        await expect(page.get_by_text('올릴 때 자동으로 고쳤습니다',exact=False).first).to_be_visible()
        await expect(page.get_by_text('증상 개선',exact=False).first).to_be_visible()
        await expect(page.get_by_text('(문장 삭제)',exact=False)).to_be_visible()

        await page.get_by_text('글자 강조 설정 바꾸기',exact=True).click()
        await page.get_by_label('핵심 문구 자동 강조').check()
        await page.get_by_label('특히 강조할 문구',exact=False).fill('핵심 기준')
        await expect(page.get_by_text('설정 저장됨',exact=True)).to_be_visible()
        assert config['phrases']==['핵심 기준'] and config['enabled'] and state['saves']>0
        await page.get_by_role('button',name='강조 미리보기',exact=True).click()
        await expect(page.get_by_label('강조 미리보기',exact=True)).to_contain_text('핵심 기준')
        OUTPUT.mkdir(parents=True,exist_ok=True)
        await page.screenshot(path=str(OUTPUT/'homepage-word-formatting.png'),full_page=True)
        await page.reload()
        await page.get_by_role('button',name='Word 원고 올리기').click()
        await page.get_by_text('글자 강조 설정 바꾸기',exact=True).click()
        await expect(page.get_by_label('핵심 문구 자동 강조')).to_be_checked()
        await expect(page.get_by_label('특히 강조할 문구',exact=False)).to_have_value('핵심 기준')
        for did in ('d1', 'd3'):
            await page.locator(f'#pick-{did}').check()
        await page.get_by_role('button',name='선택한 2개 원고 예약하기').click()
        # 이미 걸린 예약을 세어 보여 주고, 그 다음부터가 기본으로 골라져 있어야 한다.
        await expect(page.get_by_label('이미 예약된 글')).to_contain_text('이미 예약된 글 1건')
        await expect(page.get_by_label('이미 예약된 글')).to_contain_text('마지막 9/23(수) 10:00')
        # 간격을 고르면 미리보기가 저절로 뜬다 — 따로 누를 단추가 없다.
        await page.get_by_role('button',name='2시간',exact=True).click()
        await expect(page.get_by_label('예약 시각 미리보기')).to_contain_text('9/23(수) 12:00')
        assert state['modes'] and state['modes'][-1]=='after_last',state['modes']
        await page.get_by_role('button',name='예약 목록 새로 읽기').click()
        await expect(page.get_by_text('실행기가 다음 차례에',exact=False)).to_be_visible()
        assert state['rescan']==1
        await page.screenshot(path=str(OUTPUT/'homepage-word-schedule.png'),full_page=True)
        assert not state['scheduled'],'Preview must not create a reservation'
        await page.get_by_role('button',name='1건 예약하기',exact=True).click()
        await expect(page.get_by_text('예약을 걸었습니다.',exact=False)).to_be_visible()
        assert state['scheduled'],'Confirmation must submit the selected manuscript'
        assert state['deleted']==['d2'],state['deleted']
        assert not errors,errors
        await browser.close()
    print('PASS: Word upload by picker and drag-and-drop, review reason and removal, formatting, interval schedule and confirmation; API fixtures only')

asyncio.run(main())
